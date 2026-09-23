# Copyright (c) 2026 Король Дмитрий. All rights reserved.
# Проект «Лицей GPT». Автор — Король Дмитрий.

import csv
import io
import logging
import re
from pathlib import Path
from typing import List

from docx import Document as DocxDocument
from openpyxl import load_workbook
from pypdf import PdfReader

logger = logging.getLogger("rag")

TEXT_EXTS = {".txt", ".md", ".log"}
PDF_EXTS = {".pdf"}
DOCX_EXTS = {".docx"}
XLSX_EXTS = {".xlsx"}
CSV_EXTS = {".csv"}

SUPPORTED_EXTS = sorted(TEXT_EXTS | PDF_EXTS | DOCX_EXTS | XLSX_EXTS | CSV_EXTS)

DATE_PATTERNS = [
    r"\b\d{1,2}[.\-/]\d{1,2}[.\-/]\d{2,4}\b",
    r"\b\d{4}[.\-/]\d{1,2}[.\-/]\d{1,2}\b",
]
TIME_PATTERN = r"\b[0-2]?\d[:.][0-5]\d\b"


def extract_meta(text: str) -> dict:
    return {
        "dates": list({d for p in DATE_PATTERNS for d in re.findall(p, text, re.I)}),
        "times": list(set(re.findall(TIME_PATTERN, text))),
        "tags": ["olympiad"] if re.search(r"олимп\w*", text, re.I) else [],
        "size": len(text),
    }


def _read_txt(path: Path) -> str:
    for enc in ("utf-8", "utf-8-sig", "cp1251", "latin-1"):
        try:
            return path.read_text(encoding=enc)
        except UnicodeDecodeError:
            continue
    return path.read_text(encoding="utf-8", errors="ignore")


def _read_pdf(path: Path) -> str:
    reader = PdfReader(str(path))
    parts: List[str] = []
    for i, page in enumerate(reader.pages):
        try:
            txt = page.extract_text() or ""
        except Exception as e:
            logger.warning("PDF стр. %d — ошибка извлечения: %s", i + 1, e)
            txt = ""
        if txt.strip():
            parts.append(txt)
    return "\n\n".join(parts)


def _read_docx(path: Path) -> str:
    doc = DocxDocument(str(path))
    parts: List[str] = []
    for p in doc.paragraphs:
        if p.text.strip():
            parts.append(p.text)
    for t_idx, table in enumerate(doc.tables):
        parts.append(f"\n[Таблица {t_idx + 1}]")
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def _read_xlsx(path: Path) -> str:
    wb = load_workbook(str(path), data_only=True, read_only=True)
    parts: List[str] = []
    for sheet_name in wb.sheetnames:
        sheet = wb[sheet_name]
        parts.append(f"\n[Лист: {sheet_name}]")
        for row in sheet.iter_rows(values_only=True):
            vals = [str(c) if c is not None else "" for c in row]
            if any(v.strip() for v in vals):
                parts.append(" | ".join(vals))
    wb.close()
    return "\n".join(parts)


def _read_csv(path: Path) -> str:
    raw = None
    for enc in ("utf-8", "utf-8-sig", "cp1251", "latin-1"):
        try:
            raw = path.read_text(encoding=enc)
            break
        except UnicodeDecodeError:
            continue
    if raw is None:
        raw = path.read_text(encoding="utf-8", errors="ignore")

    sample = raw[:4096]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel

    reader = csv.reader(io.StringIO(raw), dialect)
    parts: List[str] = []
    for row in reader:
        if any(cell.strip() for cell in row):
            parts.append(" | ".join(cell.strip() for cell in row))
    return "\n".join(parts)


def extract_text_from_file(path: str, mime: str = "") -> str:
    p = Path(path)
    suffix = p.suffix.lower()

    try:
        if suffix in TEXT_EXTS:
            return _read_txt(p)
        if suffix in PDF_EXTS:
            return _read_pdf(p)
        if suffix in DOCX_EXTS:
            return _read_docx(p)
        if suffix in XLSX_EXTS:
            return _read_xlsx(p)
        if suffix in CSV_EXTS:
            return _read_csv(p)
        logger.warning("Неизвестное расширение %s — читаю как текст", suffix)
        return _read_txt(p)
    except Exception:
        logger.exception("Не удалось извлечь текст из файла %s", path)
        return ""


def ingest_text(db, assistant, title: str, content: str = None, file_path: str = None, owner_id=None) -> List[int]:
    if file_path:
        content = extract_text_from_file(file_path)
    if not content or not content.strip():
        raise ValueError("В файле не найден текст для индексации")

    meta = extract_meta(content)
    doc_id = db.add_doc(title, content, embedding=None, meta=meta, owner_id=owner_id)
    return [doc_id]